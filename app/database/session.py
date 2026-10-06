import logging
from datetime import date
from pathlib import Path

from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker

BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / "database" / "lonsystem.db"

engine = create_engine(
    f"sqlite:///{DB_PATH}",
    connect_args={"check_same_thread": False},
)


@event.listens_for(engine, "connect")
def set_wal_mode(dbapi_connection, connection_record):
    # pysqlite's eget implicitte transaktions-styring (den "legacy"
    # autocommit-lignende adfærd der selv udsteder BEGIN ved skrivning)
    # forhindrer SQLAlchemys SAVEPOINT (Session.begin_nested()) i at virke
    # pålideligt – en ændring lavet inde i en begin_nested()-blok kan blive
    # skrevet til databasefilen selvom den efterfølgende rulles tilbage med
    # Session.rollback() (bekræftet 2026-09-04 ved en direkte reproduktion:
    # et enkelt felt ændret i en begin_nested()-blok og derefter rullet
    # tilbage forblev ændret i en helt frisk session bagefter). Dette er den
    # velkendte, dokumenterede SQLAlchemy+pysqlite-adfærd ift. savepoints –
    # løsningen er at slå pysqlites egen transaktionsstyring fra
    # (isolation_level=None) og lade SQLAlchemy selv udstede BEGIN (se
    # "begin"-eventet nedenfor).
    dbapi_connection.isolation_level = None
    try:
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA journal_mode=WAL")
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()
    except Exception as e:
        logging.error(f"WAL-mode opsætning fejlede: {e}")


@event.listens_for(engine, "begin")
def do_begin(conn):
    conn.exec_driver_sql("BEGIN")


SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def get_db():
    db = SessionLocal()
    try:
        yield db
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def init_db():
    from database.models import Base
    Base.metadata.create_all(bind=engine)
    _migrate()
    _seed_roles()
    _seed_admin()
    _seed_master_data()
    _seed_cvr()
    _seed_holidays()
    _ensure_sh_pay_types()    # SH-løntypekoder kode 4 og 63
    _ensure_anciennitet_alert_permission()
    _ensure_paragraf_56_alert_permission()
    _ensure_manage_baselines_permission()
    _ensure_activity_permissions()
    _ensure_auto_approve_permission()
    _ensure_manage_auto_approval_permission()
    _ensure_system_settings()
    _ensure_vagtplan_permissions()
    _ensure_employee_supplements_permission()
    _ensure_toggle_springer_permission()
    _ensure_payroll_settlement_permissions()
    _ensure_edit_activities_permission()
    _ensure_elev_wage_approve_permission()
    _seed_agreement_rates_2025_2028()
    _ensure_springer_pay_type()
    _ensure_feriefri_fuldloennet_pay_type()
    _ensure_loen_andet_sted_fra_absence_type()
    _ensure_eksport_absence_type()
    _migrate_dispatcher_groups()
    _migrate_dispatcher_group_to_single()


def _migrate():
    """Tilføjer nye kolonner til eksisterende databaser uden at miste data."""
    import sqlite3 as _sqlite3
    with _sqlite3.connect(str(DB_PATH)) as conn:
        existing = {row[1] for row in conn.execute("PRAGMA table_info(activities)")}
        if "created_by" not in existing:
            conn.execute("ALTER TABLE activities ADD COLUMN created_by VARCHAR")
            conn.commit()
        ot_cols = {row[1] for row in conn.execute("PRAGMA table_info(master_overtime_rates)")}
        if "is_user_created" not in ot_cols:
            conn.execute("ALTER TABLE master_overtime_rates ADD COLUMN is_user_created BOOLEAN DEFAULT 0")
            conn.commit()
        sup_cols = {row[1] for row in conn.execute("PRAGMA table_info(master_supplement_rates)")}
        if "is_user_created" not in sup_cols:
            conn.execute("ALTER TABLE master_supplement_rates ADD COLUMN is_user_created BOOLEAN DEFAULT 0")
            conn.commit()
        pt_cols = {row[1] for row in conn.execute("PRAGMA table_info(master_pay_types)")}
        if "is_user_created" not in pt_cols:
            conn.execute("ALTER TABLE master_pay_types ADD COLUMN is_user_created BOOLEAN DEFAULT 0")
            conn.commit()
        if "csv_quantity_type" not in pt_cols:
            conn.execute("ALTER TABLE master_pay_types ADD COLUMN csv_quantity_type VARCHAR(20) NOT NULL DEFAULT 'hours'")
            conn.execute("ALTER TABLE master_pay_types ADD COLUMN csv_rate_source VARCHAR(30) NOT NULL DEFAULT 'hourly'")
            conn.execute("ALTER TABLE master_pay_types ADD COLUMN csv_include_rate BOOLEAN NOT NULL DEFAULT 1")
            conn.execute("ALTER TABLE master_pay_types ADD COLUMN csv_include_total BOOLEAN NOT NULL DEFAULT 0")
            conn.commit()
            conn.execute("UPDATE master_pay_types SET csv_rate_source='ot_before' WHERE code_key='OT_BEFORE'")
            conn.execute("UPDATE master_pay_types SET csv_rate_source='ot_13' WHERE code_key='OT_13'")
            conn.execute("UPDATE master_pay_types SET csv_rate_source='ot_extra' WHERE code_key='OT_EXTRA'")
            conn.execute("UPDATE master_pay_types SET csv_rate_source='salt' WHERE code_key='SALT'")
            conn.execute("UPDATE master_pay_types SET csv_quantity_type='count', csv_rate_source='overnight' WHERE code_key='OVERNATNING'")
            conn.execute("UPDATE master_pay_types SET csv_rate_source='dagpenge' WHERE code_key='PARAGRAF_56'")
            conn.execute("UPDATE master_pay_types SET csv_rate_source='dagpenge' WHERE code_key='BARN_1SYGEDAG'")
            conn.commit()
        if "csv_include_rate" not in pt_cols:
            conn.execute("ALTER TABLE master_pay_types ADD COLUMN csv_include_rate BOOLEAN NOT NULL DEFAULT 1")
            conn.commit()
        emp_cols = {row[1] for row in conn.execute("PRAGMA table_info(employees)")}
        if "cvr_number" not in emp_cols:
            conn.execute("ALTER TABLE employees ADD COLUMN cvr_number VARCHAR(20)")
            conn.commit()
        if "anciennitet_dismissed_at" not in emp_cols:
            conn.execute("ALTER TABLE employees ADD COLUMN anciennitet_dismissed_at DATETIME")
            conn.commit()
        if "terminsdato" not in emp_cols:
            conn.execute("ALTER TABLE employees ADD COLUMN terminsdato DATE")
            conn.commit()
        if "initials" not in emp_cols:
            conn.execute("ALTER TABLE employees ADD COLUMN initials VARCHAR(10)")
            conn.commit()
        if "dispatcher_group_id" not in emp_cols:
            conn.execute("ALTER TABLE employees ADD COLUMN dispatcher_group_id INTEGER")
            conn.commit()
        if "paragraf_56" not in emp_cols:
            conn.execute("ALTER TABLE employees ADD COLUMN paragraf_56 BOOLEAN NOT NULL DEFAULT 0")
            conn.execute("ALTER TABLE employees ADD COLUMN paragraf_56_start_date DATE")
            conn.execute("ALTER TABLE employees ADD COLUMN paragraf_56_end_date DATE")
            conn.commit()
        if "afloeser" not in emp_cols:
            conn.execute("ALTER TABLE employees ADD COLUMN afloeser BOOLEAN NOT NULL DEFAULT 0")
            conn.commit()
        act_cols2 = {row[1] for row in conn.execute("PRAGMA table_info(activities)")}
        if "deactivated_by" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN deactivated_by VARCHAR")
            conn.commit()
        if "updated_by" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN updated_by VARCHAR")
            conn.commit()
        if "auto_approved" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN auto_approved BOOLEAN NOT NULL DEFAULT 0")
            conn.commit()
        if "auto_approval_flags" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN auto_approval_flags TEXT NOT NULL DEFAULT '[]'")
            conn.commit()
        if "is_likely_incomplete" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN is_likely_incomplete BOOLEAN NOT NULL DEFAULT 0")
            conn.commit()
        if "baseline_duration_minutes" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN baseline_duration_minutes REAL")
            conn.commit()
        if "baseline_start_hour" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN baseline_start_hour REAL")
            conn.commit()
        if "hidden_from_vagtplan" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN hidden_from_vagtplan BOOLEAN NOT NULL DEFAULT 0")
            conn.commit()
        if "absence_group_id" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN absence_group_id VARCHAR(36)")
            conn.commit()
        if "original_pause_intervals" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN original_pause_intervals TEXT")
            conn.commit()
        if "original_segments" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN original_segments TEXT")
            conn.commit()
        if "vehicle_uses" not in act_cols2:
            conn.execute("ALTER TABLE activities ADD COLUMN vehicle_uses TEXT")
            conn.commit()
        existing_indexes = {row[1] for row in conn.execute("PRAGMA index_list(activities)")}
        if "ix_activities_employee_start_source" not in existing_indexes:
            conn.execute(
                "CREATE INDEX ix_activities_employee_start_source "
                "ON activities(employee_id, start_time, source)"
            )
            conn.commit()
        tachograph_index_sql = conn.execute(
            "SELECT sql FROM sqlite_master WHERE type='index' "
            "AND name='uq_activities_employee_start_tachograph'"
        ).fetchone()
        if tachograph_index_sql is not None:
            # Fjernet 2026-09-07: indekset eksisterede udelukkende for at
            # beskytte mod to SAMTIDIGE importer af samme vagt – men importen
            # kører kun for én medarbejder ad gangen, aldrig samtidig (heller
            # ikke efter den planlagte automatisering), så det scenarie kan
            # reelt ikke opstå. At en fuldstændig identisk vagt ikke
            # importeres to gange sikres i stedet fuldt ud af
            # applikationslogikken i import_ddd.py::_import_activity.
            conn.execute("DROP INDEX uq_activities_employee_start_tachograph")
            conn.commit()
        if "ix_activities_period_status" not in existing_indexes:
            conn.execute(
                "CREATE INDEX ix_activities_period_status ON activities(pay_period_id, status)"
            )
            conn.commit()
        if "ix_activities_absence_group" not in existing_indexes:
            conn.execute(
                "CREATE INDEX ix_activities_absence_group ON activities(absence_group_id)"
            )
            conn.commit()
        baseline_indexes = {row[1] for row in conn.execute("PRAGMA index_list(employee_baselines)")}
        if "uq_employee_baselines_employee_weekday" not in baseline_indexes:
            # Fjern evt. eksisterende dubletter (samme medarbejder+ugedag,
            # opstået pga. den tidligere manglende spærre) før det unikke
            # indeks oprettes. Beholder rækken med flest samples – den har
            # det mest velfunderede datagrundlag af de to.
            conn.execute(
                "DELETE FROM employee_baselines WHERE id NOT IN ("
                "  SELECT id FROM ("
                "    SELECT id, ROW_NUMBER() OVER ("
                "      PARTITION BY employee_id, weekday ORDER BY sample_count DESC, id ASC"
                "    ) AS rn FROM employee_baselines"
                "  ) WHERE rn = 1"
                ")"
            )
            conn.execute(
                "CREATE UNIQUE INDEX uq_employee_baselines_employee_weekday "
                "ON employee_baselines(employee_id, weekday)"
            )
            conn.commit()
        audit_indexes = {row[1] for row in conn.execute("PRAGMA index_list(audit_logs)")}
        if "ix_audit_logs_timestamp" not in audit_indexes:
            conn.execute("CREATE INDEX ix_audit_logs_timestamp ON audit_logs(timestamp)")
            conn.commit()
        dg_cols = {row[1] for row in conn.execute("PRAGMA table_info(dispatcher_groups)")}
        if "vehicle_id" not in dg_cols:
            conn.execute("ALTER TABLE dispatcher_groups ADD COLUMN vehicle_id INTEGER")
            conn.commit()
        if "visible_in_activity_overview" not in dg_cols:
            conn.execute(
                "ALTER TABLE dispatcher_groups ADD COLUMN visible_in_activity_overview "
                "BOOLEAN NOT NULL DEFAULT 1"
            )
            conn.commit()
        veh_cols = {row[1] for row in conn.execute("PRAGMA table_info(vehicles)")}
        if "description" not in veh_cols:
            conn.execute("ALTER TABLE vehicles ADD COLUMN description TEXT")
            conn.commit()
        if "dispatcher_group_id" not in veh_cols:
            conn.execute("ALTER TABLE vehicles ADD COLUMN dispatcher_group_id INTEGER")
            conn.commit()
        if "vognpark" not in veh_cols:
            conn.execute("ALTER TABLE vehicles ADD COLUMN vognpark BOOLEAN NOT NULL DEFAULT 0")
            conn.commit()
        emp_cols3 = {row[1] for row in conn.execute("PRAGMA table_info(employees)")}
        if "fast_bil" not in emp_cols3:
            conn.execute("ALTER TABLE employees ADD COLUMN fast_bil BOOLEAN NOT NULL DEFAULT 0")
            conn.commit()
        if "fast_bil_vehicle_id" not in emp_cols3:
            conn.execute("ALTER TABLE employees ADD COLUMN fast_bil_vehicle_id INTEGER")
            conn.commit()
        if "ot_extra_alle_timer" not in emp_cols3:
            conn.execute("ALTER TABLE employees ADD COLUMN ot_extra_alle_timer BOOLEAN NOT NULL DEFAULT 0")
            conn.commit()
        if "absence_vehicle_id" not in emp_cols3:
            conn.execute("ALTER TABLE employees ADD COLUMN absence_vehicle_id INTEGER")
            conn.commit()
        for col, ddl in (
            ("position_id", "INTEGER"),
            ("seniority_date", "DATE"),
            ("cpr_number", "VARCHAR(11)"),
            ("elev", "BOOLEAN NOT NULL DEFAULT 0"),
            ("elev_start_date", "DATE"),
            ("elev_end_date", "DATE"),
            ("voksenelev", "BOOLEAN NOT NULL DEFAULT 0"),
            # Eksisterende medarbejdere sættes bevidst IKKE som medlem (besluttet 2026-10-01);
            # nye medarbejdere får modellens default True.
            ("personaleforening", "BOOLEAN NOT NULL DEFAULT 0"),
            ("natarbejde_tillaeg", "BOOLEAN NOT NULL DEFAULT 0"),
        ):
            if col not in emp_cols3:
                conn.execute(f"ALTER TABLE employees ADD COLUMN {col} {ddl}")
                conn.commit()
        # Flere aktive tillæg pr. medarbejder er tilladt siden 2026-10-05 (summeres).
        conn.execute("DROP INDEX IF EXISTS uq_employee_supplements_one_open_row")
        conn.execute(
            "CREATE INDEX IF NOT EXISTS ix_employee_supplements_employee_id "
            "ON employee_supplements(employee_id)"
        )
        conn.commit()
        rate_cols = {row[1] for row in conn.execute("PRAGMA table_info(master_agreement_type_rates)")}
        if rate_cols and "note" not in rate_cols:
            conn.execute("ALTER TABLE master_agreement_type_rates ADD COLUMN note TEXT")
            conn.commit()
        sup_cols2 = {row[1] for row in conn.execute("PRAGMA table_info(employee_supplements)")}
        if "deactivated_at" not in sup_cols2:
            conn.execute("ALTER TABLE employee_supplements ADD COLUMN deactivated_at DATETIME")
            # Tillæg der allerede er afsluttet via "Afslut" (før 2026-10-05 gjaldt de
            # perioden ud) markeres som afsluttet; sidste gyldige dag = afslutningsdagen.
            conn.execute(
                "UPDATE employee_supplements SET "
                "deactivated_at = (SELECT MAX(a.timestamp) FROM audit_logs a "
                "  WHERE a.action = 'employee_supplement_end' AND a.entity_id = employee_supplements.id), "
                "end_date = MAX(start_date, (SELECT date(MAX(a.timestamp), 'localtime') FROM audit_logs a "
                "  WHERE a.action = 'employee_supplement_end' AND a.entity_id = employee_supplements.id)) "
                "WHERE id IN (SELECT entity_id FROM audit_logs WHERE action = 'employee_supplement_end')"
            )
            conn.commit()

        # Migrer eksisterende faste sats-kilde-værdier til det nye id-baserede skema
        # (overtime:<id> / supplement:<id>) – idempotent, rammer kun rækker der
        # stadig har en af de gamle faste værdier.
        _legacy_rate_src_map = {
            "ot_before": ("master_overtime_rates", "Overtid 1 time før"),
            "ot_13": ("master_overtime_rates", "Overtid 1-3 timer efter"),
            "ot_extra": ("master_overtime_rates", "Øvrigt overtid"),
            "salt": ("master_supplement_rates", "Salttillæg"),
            "overnight": ("master_supplement_rates", "Overnatning"),
            "dagpenge": ("master_supplement_rates", "Dagpenge §56"),
            "springer": ("master_supplement_rates", "Springertillæg"),
        }
        for old_value, (table, label) in _legacy_rate_src_map.items():
            found = conn.execute(f"SELECT id FROM {table} WHERE label = ?", (label,)).fetchone()
            if found:
                prefix = "overtime" if table == "master_overtime_rates" else "supplement"
                conn.execute(
                    "UPDATE master_pay_types SET csv_rate_source = ? WHERE csv_rate_source = ?",
                    (f"{prefix}:{found[0]}", old_value),
                )
        conn.commit()


def _seed_roles():
    from database.models import Role
    db = SessionLocal()
    try:
        if db.query(Role).count() == 0:
            for r in [
                Role(name="admin", display_name="Administrator", is_system=True,
                     permissions=["payroll", "import_ddd", "user_management", "reopen_period", "manage_baselines", "manage_auto_approval", "approve_activities", "view_calendar"]),
                Role(name="lonbogholder", display_name="Lønbogholder", is_system=False,
                     permissions=["payroll", "absence_overview", "import_ddd", "anciennitet_alert", "approve_activities", "view_calendar", "auto_approve_manual_activities"]),
                Role(name="disponent", display_name="Disponent", is_system=False,
                     permissions=["approve_activities", "view_calendar"]),
            ]:
                db.add(r)
            db.commit()
    finally:
        db.close()


def _seed_admin():
    from database.models import AppUser
    from auth import hash_password
    db = SessionLocal()
    try:
        if db.query(AppUser).count() == 0:
            db.add(AppUser(
                name="Administrator",
                initials="admin",
                email="",
                role="admin",
                password_hash=hash_password("admin"),
                active=True,
            ))
            db.commit()
    finally:
        db.close()


def _normalize_absence_key(label: str) -> str:
    overrides = {"Kursus/Skole": "skole_kursus"}
    if label in overrides:
        return overrides[label]
    s = label.lower()
    s = s.replace("§", "paragraf_")
    s = s.replace("æ", "ae").replace("ø", "oe").replace("å", "aa")
    s = s.replace(" ", "_").replace("/", "_").replace(".", "").replace("-", "_")
    while "__" in s:
        s = s.replace("__", "_")
    return s.strip("_")


def _seed_agreement_kinds(db):
    from database.models import MasterAgreementKind
    if db.query(MasterAgreementKind).count() == 0:
        db.add(MasterAgreementKind(
            key="hourly_fixed", label="Timelønnet, fast arbejdstid",
            is_active=True, is_user_created=False,
            requires_agreement_type=True, sort_order=1,
        ))
        db.add(MasterAgreementKind(
            key="hourly_flexible", label="Timelønnet, ikke fastlagt arbejdstid",
            is_active=True, is_user_created=False,
            requires_agreement_type=True, sort_order=2,
        ))
        db.commit()


def _seed_master_data():
    from decimal import Decimal
    from database.models import (
        MasterAgreementType, MasterOvertimeRate,
        MasterSupplementRate, MasterPayType, MasterAbsenceType,
        MasterAgreementKind,
    )
    db = SessionLocal()
    try:
        if db.query(MasterAgreementType).count() == 0:
            try:
                from calculators.rates_loader import load_agreement_types
                for name, rate in load_agreement_types().items():
                    db.add(MasterAgreementType(name=name, hourly_rate=rate))
                db.commit()
            except Exception as e:
                db.rollback()
                logging.warning(f"Overenskomsttyper – seeding fra Excel slog fejl: {e}")

        if db.query(MasterOvertimeRate).count() == 0:
            try:
                from calculators.rates_loader import load_overtime_rates
                for label, rate in load_overtime_rates().items():
                    db.add(MasterOvertimeRate(label=label, rate=rate))
                db.commit()
            except Exception as e:
                db.rollback()
                logging.warning(f"Overtidssatser – seeding fra Excel slog fejl: {e}")

        if db.query(MasterSupplementRate).count() == 0:
            try:
                from calculators.rates_loader import load_salt_supplement_rate, load_overnight_rate
                db.add(MasterSupplementRate(label="Salttillæg", rate=load_salt_supplement_rate()))
                db.add(MasterSupplementRate(label="Overnatning", rate=load_overnight_rate()))
                db.add(MasterSupplementRate(label="Dagpenge §56", rate=Decimal("137.43")))
                db.commit()
            except Exception as e:
                db.rollback()
                logging.warning(f"Tillægssatser – seeding slog fejl: {e}")

        if db.query(MasterPayType).count() == 0:
            from calculators.pay_rates import (
                DANLOEN_CODE_NORMAL, DANLOEN_CODE_OT_BEFORE, DANLOEN_CODE_OT_13,
                DANLOEN_CODE_OT_EXTRA, DANLOEN_CODE_SALT, DANLOEN_CODE_AFSPADSERING,
                DANLOEN_CODE_SYGDOM, DANLOEN_CODE_FERIEFRI, DANLOEN_CODE_BARSEL,
                DANLOEN_CODE_SKOLE_KURSUS, DANLOEN_CODE_OVERNATNING,
                DANLOEN_CODE_PARAGRAF_56, DANLOEN_CODE_BARN_1SYGEDAG,
            )
            pay_types = [
                # (code_key, label, danloen_code, include_in_csv, sort_order, qty_type, rate_src)
                ("NORMAL",        "Normal tid",           DANLOEN_CODE_NORMAL,        True,  1, "hours", "hourly"),
                ("OT_BEFORE",     "Overtid 1 time før",   DANLOEN_CODE_OT_BEFORE,     True,  2, "hours", "ot_before"),
                ("OT_13",         "Overtid 1-3 timer",    DANLOEN_CODE_OT_13,         True,  3, "hours", "ot_13"),
                ("OT_EXTRA",      "Øvrig overtid",        DANLOEN_CODE_OT_EXTRA,      True,  4, "hours", "ot_extra"),
                ("SALT",          "Salttillæg",           DANLOEN_CODE_SALT,          True,  5, "hours", "salt"),
                ("OVERNATNING",   "Overnatning",          DANLOEN_CODE_OVERNATNING,   True,  6, "count", "overnight"),
                ("AFSPADSERING",  "Afspadsering",         DANLOEN_CODE_AFSPADSERING,  True,  7, "hours", "hourly"),
                ("SYGDOM",        "Sygdom med løn",       DANLOEN_CODE_SYGDOM,        True,  8, "hours", "hourly"),
                ("PARAGRAF_56",   "§56 syg",              DANLOEN_CODE_PARAGRAF_56,   True,  9, "hours", "dagpenge"),
                ("BARN_1SYGEDAG", "Barn 1.sygedag",       DANLOEN_CODE_BARN_1SYGEDAG, True, 10, "hours", "dagpenge"),
                ("FERIEFRI",      "Feriefri",             DANLOEN_CODE_FERIEFRI,      True, 11, "hours", "hourly"),
                ("BARSEL",        "Barsel",               DANLOEN_CODE_BARSEL,        True, 12, "hours", "hourly"),
                ("SKOLE_KURSUS",  "Kursus/Skole",         DANLOEN_CODE_SKOLE_KURSUS,  True, 13, "hours", "hourly"),
            ]
            for ck, lbl, code, inc, order, qty, rate_src in pay_types:
                db.add(MasterPayType(
                    code_key=ck, label=lbl, danloen_code=code,
                    include_in_csv=inc, sort_order=order,
                    csv_quantity_type=qty, csv_rate_source=rate_src,
                    csv_include_rate=True, csv_include_total=False,
                ))
            db.commit()
        if db.query(MasterAbsenceType).count() == 0:
            try:
                from openpyxl import load_workbook
                xlsx_path = BASE_DIR / "Fraværstyper.xlsx"
                _backend_only = {
                    "sygdom_u_8uger", "sygdom_u_8_uger",
                    "barn_1sygedag_u_8uger", "barsel_u_loen",
                }
                if xlsx_path.exists():
                    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
                    ws = wb.active
                    order = 1
                    first = True
                    for row in ws.iter_rows(values_only=True):
                        if first:
                            first = False
                            continue
                        cell = row[0]
                        if cell:
                            lbl = str(cell).strip()
                            key = _normalize_absence_key(lbl)
                            if key not in _backend_only:
                                db.add(MasterAbsenceType(
                                    label=lbl, normalized_key=key,
                                    is_active=True, is_user_created=False,
                                    sort_order=order,
                                ))
                                order += 1
                    wb.close()
                    db.commit()
            except Exception as e:
                db.rollback()
                logging.warning(f"Fraværstyper – seeding fra Excel slog fejl: {e}")

        _seed_agreement_kinds(db)

    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved seeding af stamdata: {e}")
    finally:
        db.close()


def _seed_cvr():
    from database.models import MasterCvrNumber
    from calculators.pay_rates import CVR_NUMBER
    db = SessionLocal()
    try:
        if db.query(MasterCvrNumber).count() == 0:
            db.add(MasterCvrNumber(
                cvr_number=CVR_NUMBER,
                company_name="Poul Schou A/S",
                is_default=True,
            ))
            db.commit()
    finally:
        db.close()


def _seed_holidays():
    from database.models import Holiday
    from calculators.holidays import get_holidays_for_year
    from datetime import date
    db = SessionLocal()
    try:
        current_year = date.today().year
        for year in range(current_year, current_year + 5):
            for h in get_holidays_for_year(year):
                if not db.query(Holiday).filter(Holiday.date == h["date"]).first():
                    db.add(Holiday(
                        date=h["date"],
                        name=h["name"],
                        half_day_from=h["half_day_from"],
                        is_auto_generated=True,
                    ))
        db.commit()
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved seeding af helligdage: {e}")
    finally:
        db.close()


def _ensure_sh_pay_types():
    """Tilføjer SH-løntypekoder til eksisterende databaser (idempotent)."""
    from database.models import MasterPayType
    from calculators.pay_rates import DANLOEN_CODE_SH_FULDLOENNET, DANLOEN_CODE_SH_TIMELOENNET
    db = SessionLocal()
    try:
        entries = [
            ("SH_FULDLOENNET", "Søgnehelligdag", DANLOEN_CODE_SH_FULDLOENNET, True, 14, "hours", "hourly"),
            ("SH_TIMELOENNET", "SH-Udbetaling", DANLOEN_CODE_SH_TIMELOENNET, True, 15, "hours", "hourly"),
        ]
        for ck, lbl, code, inc, order, qty, rate_src in entries:
            existing = db.query(MasterPayType).filter(MasterPayType.code_key == ck).first()
            if not existing:
                db.add(MasterPayType(
                    code_key=ck, label=lbl, danloen_code=code,
                    include_in_csv=inc, sort_order=order,
                    csv_quantity_type=qty, csv_rate_source=rate_src,
                    csv_include_rate=True, csv_include_total=False,
                ))
            elif existing.label != lbl:
                existing.label = lbl
        db.commit()
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved seeding af SH-løntypekoder: {e}")
    finally:
        db.close()


def _ensure_springer_pay_type():
    """Seeder Springertillæg-sats og -løntypekode til eksisterende databaser (idempotent)."""
    from decimal import Decimal
    from database.models import MasterSupplementRate, MasterPayType
    from calculators.pay_rates import DANLOEN_CODE_SPRINGERTILLAEG
    db = SessionLocal()
    try:
        if not db.query(MasterSupplementRate).filter(MasterSupplementRate.label == "Springertillæg").first():
            db.add(MasterSupplementRate(label="Springertillæg", rate=Decimal("0")))
        if not db.query(MasterPayType).filter(MasterPayType.code_key == "SPRINGERTILLAEG").first():
            db.add(MasterPayType(
                code_key="SPRINGERTILLAEG", label="Springertillæg",
                danloen_code=DANLOEN_CODE_SPRINGERTILLAEG,
                include_in_csv=True, sort_order=16,
                csv_quantity_type="hours", csv_rate_source="springer",
                csv_include_rate=True, csv_include_total=False,
            ))
        db.commit()
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved seeding af Springertillæg: {e}")
    finally:
        db.close()


def _ensure_feriefri_fuldloennet_pay_type():
    """Seeder FERIEFRI_FULDLOENNET-løntypekode til eksisterende databaser (idempotent)."""
    from database.models import MasterPayType
    from calculators.pay_rates import DANLOEN_CODE_FERIEFRI_FULDLOENNET
    db = SessionLocal()
    try:
        if not db.query(MasterPayType).filter(MasterPayType.code_key == "FERIEFRI_FULDLOENNET").first():
            db.add(MasterPayType(
                code_key="FERIEFRI_FULDLOENNET", label="Feriefri fuldlønnet",
                danloen_code=DANLOEN_CODE_FERIEFRI_FULDLOENNET,
                include_in_csv=True, sort_order=17,
                csv_quantity_type="hours", csv_rate_source="hourly",
                csv_include_rate=True, csv_include_total=False,
            ))
        db.commit()
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved seeding af FERIEFRI_FULDLOENNET-løntypekode: {e}")
    finally:
        db.close()


def _ensure_loen_andet_sted_fra_absence_type():
    """Seeder fraværstypen 'Løn andet sted fra' til eksisterende databaser
    (idempotent) – opfører sig som Selvbetalt fridag (ingen etableret
    betalingsregel i _ABSENCE_LABELS, se payroll_router.py: 0 kr., ingen
    CSV-linje)."""
    from database.models import MasterAbsenceType
    db = SessionLocal()
    try:
        if not db.query(MasterAbsenceType).filter(
            MasterAbsenceType.normalized_key == "loen_andet_sted_fra"
        ).first():
            max_order = db.query(MasterAbsenceType).count()
            db.add(MasterAbsenceType(
                label="Løn andet sted fra", normalized_key="loen_andet_sted_fra",
                is_active=True, is_user_created=False,
                sort_order=max_order + 1,
            ))
        db.commit()
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved seeding af fraværstypen 'Løn andet sted fra': {e}")
    finally:
        db.close()


def _ensure_eksport_absence_type():
    """Seeder fraværstypen 'Eksport' til eksisterende databaser (idempotent) –
    opfører sig præcis som 'Løn andet sted fra' (0 kr., ingen CSV-linje)."""
    from database.models import MasterAbsenceType
    db = SessionLocal()
    try:
        if not db.query(MasterAbsenceType).filter(
            MasterAbsenceType.normalized_key == "eksport"
        ).first():
            max_order = db.query(MasterAbsenceType).count()
            db.add(MasterAbsenceType(
                label="Eksport", normalized_key="eksport",
                is_active=True, is_user_created=False,
                sort_order=max_order + 1,
            ))
        db.commit()
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved seeding af fraværstypen 'Eksport': {e}")
    finally:
        db.close()


def _grant_permissions_once(key: str, perms: list[str], role_name: str | None = None,
                            lonbogholder_only: bool = False):
    """Tildeler `perms` til rollerne ÉN gang pr. database (husket i
    applied_permission_grants under `key`). Uden denne spærre blev tildelingen
    gentaget ved hver serverstart, så en rettighed en administrator bevidst havde
    fjernet fra en rolle kom tilbage ved næste (natlige) genstart.

    role_name=None → alle roller. lonbogholder_only → kun lonbogholder, og kun hvis
    den ikke er en systemrolle (samme betingelse som de oprindelige funktioner)."""
    from database.models import AppliedPermissionGrant, Role
    db = SessionLocal()
    try:
        if db.query(AppliedPermissionGrant).filter(AppliedPermissionGrant.key == key).first():
            return
        q = db.query(Role)
        if lonbogholder_only:
            q = q.filter(Role.name == "lonbogholder", Role.is_system == False)  # noqa: E712
        elif role_name:
            q = q.filter(Role.name == role_name)
        for role in q.all():
            current = list(role.permissions or [])
            missing = [p for p in perms if p not in current]
            if missing:
                role.permissions = current + missing
        db.add(AppliedPermissionGrant(key=key))
        db.commit()
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved tildeling af rettigheder ({key}): {e}")
    finally:
        db.close()


def _ensure_anciennitet_alert_permission():
    """Tilføjer anciennitet_alert til lonbogholder-rollen (kun én gang)."""
    _grant_permissions_once("anciennitet_alert", ["anciennitet_alert"], lonbogholder_only=True)


def _ensure_paragraf_56_alert_permission():
    """Tilføjer paragraf_56_alert til lonbogholder-rollen (kun én gang)."""
    _grant_permissions_once("paragraf_56_alert", ["paragraf_56_alert"], lonbogholder_only=True)


def _migrate_dispatcher_groups():
    """
    Migrerer eksisterende medarbejderes (legacy) dispatcher_group-streng til
    den nye many-to-many-tabel. Idempotent – dropper legacy-kolonnen efter
    migrering. Opretter ingen faste/hardcodede grupper; grupper kommer
    udelukkende fra faktiske medarbejderdata eller senere CRUD via UI.
    """
    import sqlite3 as _sqlite3
    from database.models import DispatcherGroup

    with _sqlite3.connect(str(DB_PATH)) as conn:
        emp_cols = {row[1] for row in conn.execute("PRAGMA table_info(employees)")}
        if "dispatcher_group" not in emp_cols:
            # Migrering allerede udført tidligere - rør ikke ved disponentgrupperne igen.
            return

    db = SessionLocal()
    try:
        with _sqlite3.connect(str(DB_PATH)) as conn:
            rows = conn.execute(
                "SELECT id, dispatcher_group FROM employees "
                "WHERE dispatcher_group IS NOT NULL AND dispatcher_group != ''"
            ).fetchall()
            groups_by_name = {g.name: g.id for g in db.query(DispatcherGroup).all()}
            for emp_id, group_name in rows:
                group_id = groups_by_name.get(group_name)
                if group_id is None:
                    g = DispatcherGroup(name=group_name)
                    db.add(g)
                    db.commit()
                    db.refresh(g)
                    groups_by_name[group_name] = g.id
                    group_id = g.id
                already_linked = conn.execute(
                    "SELECT 1 FROM employee_dispatcher_groups WHERE employee_id = ? AND dispatcher_group_id = ?",
                    (emp_id, group_id),
                ).fetchone()
                if not already_linked:
                    conn.execute(
                        "INSERT INTO employee_dispatcher_groups (employee_id, dispatcher_group_id) VALUES (?, ?)",
                        (emp_id, group_id),
                    )
            conn.commit()
            conn.execute("ALTER TABLE employees DROP COLUMN dispatcher_group")
            conn.commit()
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved migrering af disponentgrupper: {e}")
    finally:
        db.close()


def _migrate_dispatcher_group_to_single():
    """
    Reducerer disponentgruppe fra mange-til-mange til én gruppe pr. medarbejder.
    Har en medarbejder i dag flere grupper, beholdes den alfabetisk først
    sorterede (efter gruppenavn). Idempotent – dropper employee_dispatcher_groups
    efter migrering; kører derfor kun data-trinnet én gang (tabellen er væk bagefter).
    Kolonnerne dispatcher_group_id/vehicle_id oprettes af _migrate(), som kører
    før denne funktion, så de findes allerede når vi når hertil.
    """
    import sqlite3 as _sqlite3

    try:
        with _sqlite3.connect(str(DB_PATH)) as conn:
            tables = {row[0] for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )}
            if "employee_dispatcher_groups" not in tables:
                return

            rows = conn.execute(
                "SELECT edg.employee_id, dg.id "
                "FROM employee_dispatcher_groups edg "
                "JOIN dispatcher_groups dg ON dg.id = edg.dispatcher_group_id "
                "ORDER BY edg.employee_id, dg.name"
            ).fetchall()
            primary_by_employee = {}
            for emp_id, group_id in rows:
                primary_by_employee.setdefault(emp_id, group_id)
            for emp_id, group_id in primary_by_employee.items():
                conn.execute(
                    "UPDATE employees SET dispatcher_group_id = ? "
                    "WHERE id = ? AND dispatcher_group_id IS NULL",
                    (group_id, emp_id),
                )
            conn.commit()
            conn.execute("DROP TABLE employee_dispatcher_groups")
            conn.commit()
    except Exception as e:
        logging.error(f"Fejl ved migrering af disponentgruppe til én-til-én: {e}")


def _ensure_manage_baselines_permission():
    """Tilføjer manage_baselines til admin-rollen (kun én gang)."""
    _grant_permissions_once("manage_baselines", ["manage_baselines"], role_name="admin")


def _ensure_manage_auto_approval_permission():
    """Tilføjer manage_auto_approval til admin-rollen (kun én gang)."""
    _grant_permissions_once("manage_auto_approval", ["manage_auto_approval"], role_name="admin")


def _ensure_system_settings():
    """Opretter singleton-recorden for systemindstillinger hvis den mangler (idempotent)."""
    from database.models import SystemSettings
    db = SessionLocal()
    try:
        if db.query(SystemSettings).filter(SystemSettings.id == 1).first() is None:
            db.add(SystemSettings(id=1, auto_approval_enabled=True))
            db.commit()
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved oprettelse af systemindstillinger: {e}")
    finally:
        db.close()


def _ensure_auto_approve_permission():
    """Tilføjer auto_approve_manual_activities til lonbogholder-rollen (kun én gang)."""
    _grant_permissions_once("auto_approve_manual_activities", ["auto_approve_manual_activities"], role_name="lonbogholder")


def _ensure_activity_permissions():
    """Tilføjer approve_activities og view_calendar til alle roller (kun én gang)."""
    _grant_permissions_once("activity_permissions", ["approve_activities", "view_calendar"])


def _ensure_vagtplan_permissions():
    """Tilføjer vagtplan_view + vagtplan_edit_all til ALLE roller (kun én gang) – 'alle
    nuværende roller skal kunne se og redigere i vagtplanen' (spec-beslutning 2026-08-21).
    vagtplan_edit_own tilføjes IKKE automatisk – det er en mere restriktiv, opt-in ret."""
    _grant_permissions_once("vagtplan_permissions", ["vagtplan_view", "vagtplan_edit_all"])


def _ensure_employee_supplements_permission():
    """Tilføjer manage_employee_supplements til lonbogholder-rollen (kun én gang)."""
    _grant_permissions_once("manage_employee_supplements", ["manage_employee_supplements"], lonbogholder_only=True)


def _ensure_toggle_springer_permission():
    """Tilføjer toggle_springer til ALLE roller (kun én gang)."""
    _grant_permissions_once("toggle_springer", ["toggle_springer"])


def _ensure_edit_activities_permission():
    """Tilføjer edit_activities til ALLE roller (kun én gang) – rettigheden blev indført
    2026-09-30; før da kunne alle indloggede oprette/rette aktiviteter, så ingen
    eksisterende rolle mister adgang ved indførelsen."""
    _grant_permissions_once("edit_activities", ["edit_activities"])


# Lærlingeoverenskomsten 2025-2028 § 8 (stk. 1, disponentspecialet og stk. 6 EGU).
_LAERLING_RATES = {
    # type: (1.5.2025, 1.3.2026, 1.3.2027)
    "Lærling (EUD) Sidste år af lærerkontrakt": ("115.14", "119.17", "123.34"),
    "Lærling (EUD) Næstsidsteår af lærerkontrakt": ("102.17", "105.75", "109.45"),
    "Lærling (EUD) Tredjesidste år af lærerkontrakt": ("90.54", "93.71", "96.99"),
    "Lærling (EUD) Disponentspecialet, sidste uddannelsestrin": ("122.16", "126.43", "130.86"),
    "EGU-elever": ("79.57", "82.36", "85.24"),
}
_RATE_DATES = (date(2025, 5, 1), date(2026, 3, 1), date(2027, 3, 1))


def _seed_agreement_rates_2025_2028():
    """Indlægger satserne fra Lærlingeoverenskomsten 2025-2028 som daterede satser – ÉN gang
    pr. database (2026-10-06). Ældre satser gemmes som historik; satser hvis dato er
    nået træder i kraft med det samme (fx EGU 70,74 → 82,36), fremtidige venter."""
    from decimal import Decimal
    from datetime import datetime as _dt
    from database.models import AppliedPermissionGrant, AuditLog, MasterAgreementType, MasterAgreementTypeRate
    key = "seed_agreement_rates_2025_2028"
    db = SessionLocal()
    try:
        if db.query(AppliedPermissionGrant).filter(AppliedPermissionGrant.key == key).first():
            return
        for name, rates in _LAERLING_RATES.items():
            t = db.query(MasterAgreementType).filter(MasterAgreementType.name == name).first()
            if t is None:
                if "Disponentspecialet" not in name:
                    # Findes typen ikke under præcis dette navn (fx omdøbt i Stamdata), oprettes
                    # der ikke en dublet – satserne kan lægges ind i hånden under Fremtidige satser.
                    logging.warning(f"Daterede satser: overenskomsttypen '{name}' findes ikke – sprunget over")
                    continue
                t = MasterAgreementType(name=name, hourly_rate=Decimal(rates[0]))
                db.add(t)
                db.flush()
            existing = {r.valid_from for r in db.query(MasterAgreementTypeRate)
                        .filter(MasterAgreementTypeRate.agreement_type_id == t.id).all()}
            current = Decimal(str(t.hourly_rate))
            # k = den sats der gælder nu (hvis den findes i tabellen): den og ældre gemmes som
            # historik; nyere satser lægges ind og træder i kraft på deres dato.
            k = next((i for i, r in enumerate(rates) if Decimal(r) == current), None)
            for i, (when, rate) in enumerate(zip(_RATE_DATES, rates)):
                if when in existing:
                    continue
                history = k is not None and i <= k
                db.add(MasterAgreementTypeRate(
                    agreement_type_id=t.id, valid_from=when, hourly_rate=Decimal(rate),
                    old_rate=Decimal(rates[i - 1]) if history and i > 0 else None,
                    applied_at=_dt.now() if history else None, created_by="SYSTEM",
                    note="Lærlingeoverenskomsten 2025-2028"))
            db.add(AuditLog(user_initials="SYSTEM", action="stamdata_create", entity_type="agreement_rate",
                            entity_id=t.id, details=f"Daterede satser indlagt (Lærlingeoverenskomsten 2025-2028): "
                            f"'{name}' " + ", ".join(f"{d}: {r}" for d, r in zip(_RATE_DATES, rates))))
        db.add(AppliedPermissionGrant(key=key))
        db.commit()
        from calculators.agreement_rates import apply_due_rates
        apply_due_rates(db, date.today())
    except Exception as e:
        db.rollback()
        logging.error(f"Fejl ved indlægning af daterede satser: {e}")
    finally:
        db.close()


def _ensure_elev_wage_approve_permission():
    """Tilføjer elev_wage_approve til lonbogholder-rollen (kun én gang, 2026-10-06)."""
    _grant_permissions_once("elev_wage_approve", ["elev_wage_approve"], lonbogholder_only=True)


def _ensure_payroll_settlement_permissions():
    """Tilføjer payroll_settlement_view + payroll_settlement_export til lonbogholder-rollen (kun én gang)."""
    _grant_permissions_once("payroll_settlement_permissions", ["payroll_settlement_view", "payroll_settlement_export"], lonbogholder_only=True)