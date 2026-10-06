import io
import logging
from datetime import date, datetime, timedelta
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from auth import get_current_user, log_action, require_any_permission, require_permission, user_has_any_permission, user_has_permission
from calculators.rates_loader import (
    load_agreement_types_from_db,
    seniority_variant_exists_from_db,
)
from database.session import get_db
from utils.employee_rules import (
    cpr_birthdate, in_alert_window, is_masked_cpr, jubilee_alert, mask_cpr,
    next_employee_number, round_birthday_alert, validate_cpr,
)
from utils.natural_sort import natural_key
from database.models import AppUser, DispatcherGroup, Employee, MasterAgreementKind, MasterPosition, Paragraf56AlertDismissal, Vehicle
from database.schemas import (
    AnciennitetsAlert,
    DispatcherGroupResponse,
    EmployeeCreate,
    EmployeeExportRequest,
    EmployeeResponse,
    EmployeeUpdate,
    MilestoneAlert,
    MilestoneAlertDismiss,
    Paragraf56Alert,
    Paragraf56AlertDismiss,
    Paragraf56AlertsResponse,
    WorkSchedule,
)

router = APIRouter(prefix="/api/employees", tags=["employees"])


def _months_employed(hire_date: date, today: date = None) -> int:
    if today is None:
        today = date.today()
    months = (today.year - hire_date.year) * 12 + (today.month - hire_date.month)
    if today.day < hire_date.day:
        months -= 1
    return max(0, months)


def _validate_paragraf_56(active: bool, start: Optional[date], end: Optional[date]) -> tuple:
    if not active:
        return None, None
    if not start or not end:
        raise HTTPException(400, "Start- og slutdato for §56 skal udfyldes")
    if end < start:
        raise HTTPException(400, "§56 slutdato skal være efter startdato")
    return start, end


def _sweep_expired_paragraf_56(db: Session) -> None:
    """Deaktiverer automatisk §56 for medarbejdere hvor slutdatoen er overskredet.
    Kører uafhængigt af paragraf_56_alert-tilladelsen (se list_employees()), så
    deaktiveringen sker uanset hvilke roller der har advarslen slået til. Datoerne
    bevares bevidst (ikke nulstillet), så de kan indgå i "udløbet"-informationen."""
    today = date.today()
    expired = db.query(Employee).filter(
        Employee.paragraf_56 == True,
        Employee.paragraf_56_end_date.isnot(None),
        Employee.paragraf_56_end_date < today,
    ).all()
    for emp in expired:
        emp.paragraf_56 = False
    if expired:
        db.commit()


def _paragraf56_alert(emp: Employee) -> Paragraf56Alert:
    return Paragraf56Alert(
        employee_id=emp.id,
        employee_name=emp.name,
        employee_number=emp.employee_number,
        paragraf_56_end_date=emp.paragraf_56_end_date,
    )


def _to_response(emp: Employee, db) -> EmployeeResponse:
    try:
        rate = float(load_agreement_types_from_db(db).get(emp.agreement_type, 0)) or None
    except Exception:
        rate = None
    return EmployeeResponse(
        id=emp.id,
        employee_number=emp.employee_number,
        tachograph_card_number=emp.tachograph_card_number,
        first_name=emp.first_name,
        last_name=emp.last_name,
        name=emp.name,
        address=emp.address,
        postal_code=emp.postal_code,
        email=emp.email,
        phone=emp.phone,
        mobile=emp.mobile,
        agreement_kind=emp.agreement_kind,
        agreement_type=emp.agreement_type,
        hourly_rate=rate,
        fuldloennet=emp.fuldloennet,
        active=emp.active,
        hire_date=emp.hire_date,
        termination_date=emp.termination_date,
        work_schedule=WorkSchedule(**emp.work_schedule),
        months_employed=_months_employed(emp.hire_date),
        dispatcher_group=DispatcherGroupResponse.model_validate(emp.dispatcher_group) if emp.dispatcher_group else None,
        cvr_number=emp.cvr_number,
        anciennitet_dismissed_at=emp.anciennitet_dismissed_at,
        terminsdato=emp.terminsdato,
        initials=emp.initials,
        paragraf_56=emp.paragraf_56,
        paragraf_56_start_date=emp.paragraf_56_start_date,
        paragraf_56_end_date=emp.paragraf_56_end_date,
        afloeser=emp.afloeser,
        ot_extra_alle_timer=emp.ot_extra_alle_timer,
        fast_bil=emp.fast_bil,
        fast_bil_vehicle_id=emp.fast_bil_vehicle_id,
        fast_bil_vehicle_number=emp.fast_bil_vehicle.vehicle_number if emp.fast_bil_vehicle else None,
        absence_vehicle_id=emp.absence_vehicle_id,
        absence_vehicle_number=emp.absence_vehicle.vehicle_number if emp.absence_vehicle else None,
        position_id=emp.position_id,
        position_name=emp.position.name if emp.position else None,
        seniority_date=emp.seniority_date,
        cpr_number=emp.cpr_number,
        elev=emp.elev,
        elev_start_date=emp.elev_start_date,
        elev_end_date=emp.elev_end_date,
        voksenelev=emp.voksenelev,
        personaleforening=emp.personaleforening,
        natarbejde_tillaeg=emp.natarbejde_tillaeg,
    )


# Medarbejderlisten (navne, grupper, skema m.m.) bruges af alle skærmbilleder der viser
# medarbejdere – men kontakt-/løn-/kortoplysninger kun med 'Se medarbejdere'.
_employee_list_access = require_any_permission(
    "view_employees", "manage_employees", "view_calendar", "edit_activities", "approve_activities",
    "vagtplan_view", "dagsplan_view", "payroll", "payroll_settlement_view", "absence_overview",
    "manage_employee_supplements", "stamdata",
)
_PRIVATE_EMPLOYEE_FIELDS = (
    "tachograph_card_number", "address", "postal_code", "email", "phone", "mobile",
    "hourly_rate", "cvr_number", "cpr_number",
)


def _apply_cpr_mask(resp: EmployeeResponse, db, current_user: AppUser) -> EmployeeResponse:
    """De sidste fire cifre i CPR sendes kun til brugere med 'Se CPR-nummer'."""
    if resp.cpr_number and not user_has_permission(db, current_user, "view_cpr"):
        resp.cpr_number = mask_cpr(resp.cpr_number)
    return resp


def _visible_response(emp: Employee, db, current_user: AppUser) -> EmployeeResponse:
    """Fuld stamdata kun med 'Se medarbejdere'/'Tilføj medarbejdere' – ellers
    udelades kontakt-, løn-, førerkort- og CPR-oplysninger. CPR maskeres desuden
    uden 'Se CPR-nummer'."""
    resp = _to_response(emp, db)
    if not user_has_any_permission(db, current_user, "view_employees", "manage_employees"):
        for field in _PRIVATE_EMPLOYEE_FIELDS:
            setattr(resp, field, None)
        return resp
    return _apply_cpr_mask(resp, db, current_user)


@router.get("", response_model=list[EmployeeResponse])
def list_employees(active_only: bool = True,
                   current_user: AppUser = Depends(_employee_list_access),
                   db: Session = Depends(get_db)):
    _sweep_expired_paragraf_56(db)
    q = db.query(Employee)
    if active_only:
        q = q.filter(Employee.active == True)
    return [_visible_response(e, db, current_user)
            for e in q.order_by(Employee.last_name, Employee.first_name).all()]


@router.get("/agreement-types")
def agreement_types(current_user: AppUser = Depends(get_current_user),
                    db: Session = Depends(get_db)):
    """Overenskomsttyper og timesatser fra stamdata-tabellen."""
    try:
        types = load_agreement_types_from_db(db)
    except Exception as e:
        logging.error(f"Overenskomsttyper kunne ikke indlæses: {e}")
        raise HTTPException(500, "Overenskomsttyper kunne ikke indlæses – kontakt administrator")
    return [{"name": k, "hourly_rate": float(v)} for k, v in types.items()]


@router.get("/agreement-kinds")
def agreement_kinds(current_user: AppUser = Depends(get_current_user),
                    db: Session = Depends(get_db)):
    """Aktive Aftale-typer fra Stamdata – bruges til at udfylde medarbejder-modalens dropdown."""
    rows = db.query(MasterAgreementKind).filter(MasterAgreementKind.is_active == True).order_by(
        MasterAgreementKind.sort_order, MasterAgreementKind.label
    ).all()
    return [
        {"key": r.key, "label": r.label, "requires_agreement_type": r.requires_agreement_type}
        for r in rows
    ]


def _agreement_type_required(db: Session, agreement_kind: str) -> bool:
    row = db.query(MasterAgreementKind).filter(MasterAgreementKind.key == agreement_kind).first()
    return row.requires_agreement_type if row else True


@router.get("/dispatcher-groups", response_model=list[DispatcherGroupResponse])
def dispatcher_groups(current_user: AppUser = Depends(get_current_user),
                      db: Session = Depends(get_db)):
    """Liste over disponentgrupper – bruges til at udfylde medarbejder-modalens afkrydsningsliste."""
    return sorted(db.query(DispatcherGroup).all(), key=lambda g: natural_key(g.name))


def _resolve_dispatcher_group(db: Session, group_id: Optional[int]) -> Optional[DispatcherGroup]:
    if group_id is None:
        return None
    group = db.query(DispatcherGroup).filter(DispatcherGroup.id == group_id).first()
    if not group:
        raise HTTPException(400, f"Ukendt disponentgruppe-id: {group_id}")
    return group


def _resolve_vehicle_id(db: Session, vehicle_id: Optional[int]) -> Optional[int]:
    if vehicle_id is None:
        return None
    if not db.query(Vehicle).filter(Vehicle.id == vehicle_id).first():
        raise HTTPException(400, f"Ukendt vogn-id: {vehicle_id}")
    return vehicle_id


FUNKTIONAER = "funktionaer"


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _validate_employee_fields(db: Session, kind: str, values: dict, check: set) -> None:
    """Påkrævede felter pr. medarbejdertype (spec afsnit 3). Kun felter i `check`
    kontrolleres – ved PATCH er det de felter, der sendes med, så gamle
    medarbejdere først skal udfylde dem næste gang hele formularen gemmes.
    Skjulte felter (fx førerkort for funktionærer) er aldrig påkrævede."""
    is_office = kind == FUNKTIONAER
    required = {"position_id": "Stilling", "email": "Email"}
    if is_office:
        required["initials"] = "Initialer"
    else:
        required["tachograph_card_number"] = "Førerkortnummer"
    missing = [label for field, label in required.items() if field in check and _blank(values.get(field))]
    if missing:
        raise HTTPException(400, f"Påkrævede felter mangler: {', '.join(missing)}")
    if "position_id" in check and values.get("position_id") is not None:
        if not db.query(MasterPosition).filter(MasterPosition.id == values["position_id"]).first():
            raise HTTPException(400, f"Ukendt stilling-id: {values['position_id']}")
    if not is_office and values.get("elev") and ({"elev", "elev_start_date", "elev_end_date"} & check):
        start, end = values.get("elev_start_date"), values.get("elev_end_date")
        if not start or not end:
            raise HTTPException(400, "Start- og slutdato for elev skal udfyldes")
        if end < start:
            raise HTTPException(400, "Elev slutdato skal være efter startdato")


def _clean_cpr(value: Optional[str]) -> Optional[str]:
    if _blank(value):
        return None
    try:
        return validate_cpr(value)
    except ValueError as e:
        raise HTTPException(400, str(e))


@router.post("", response_model=EmployeeResponse, status_code=201)
def create_employee(body: EmployeeCreate,
                    current_user: AppUser = Depends(require_permission("manage_employees")),
                    db: Session = Depends(get_db)):
    if db.query(Employee).filter(Employee.employee_number == body.employee_number).first():
        raise HTTPException(400, "Lønnummer eksisterer allerede")
    if body.tachograph_card_number:
        if db.query(Employee).filter(Employee.tachograph_card_number == body.tachograph_card_number).first():
            raise HTTPException(400, "Førerkortnummer eksisterer allerede")
    if not db.query(MasterAgreementKind).filter(MasterAgreementKind.key == body.agreement_kind).first():
        raise HTTPException(400, f"Ukendt aftaletype: {body.agreement_kind}")
    if _agreement_type_required(db, body.agreement_kind):
        if not body.agreement_type or body.agreement_type not in load_agreement_types_from_db(db):
            raise HTTPException(400, f"Ukendt overenskomsttype: {body.agreement_type}")
    else:
        body.agreement_type = ""
    body.paragraf_56_start_date, body.paragraf_56_end_date = _validate_paragraf_56(
        body.paragraf_56, body.paragraf_56_start_date, body.paragraf_56_end_date
    )
    body.cpr_number = _clean_cpr(body.cpr_number)
    _validate_employee_fields(db, body.agreement_kind, body.model_dump(), set(EmployeeCreate.model_fields))

    data = body.model_dump(exclude={"dispatcher_group_id", "fast_bil_vehicle_id", "absence_vehicle_id"})
    data["work_schedule"] = body.work_schedule.model_dump()
    data["fast_bil_vehicle_id"] = _resolve_vehicle_id(db, body.fast_bil_vehicle_id)
    data["absence_vehicle_id"] = _resolve_vehicle_id(db, body.absence_vehicle_id)
    emp = Employee(**data)
    emp.dispatcher_group = _resolve_dispatcher_group(db, body.dispatcher_group_id)
    db.add(emp)
    db.commit()
    db.refresh(emp)
    return _apply_cpr_mask(_to_response(emp, db), db, current_user)


@router.get("/next-employee-number")
def next_number(current_user: AppUser = Depends(require_permission("manage_employees")),
                db: Session = Depends(get_db)):
    """Forslag til lønnummer: højeste numeriske lønnummer >= 34000 plus 1 (aktive og inaktive)."""
    numbers = [n for (n,) in db.query(Employee.employee_number).all()]
    return {"suggestion": next_employee_number(numbers)}


@router.get("/check-number")
def check_number(number: str, exclude_id: Optional[int] = None,
                 current_user: AppUser = Depends(require_permission("manage_employees")),
                 db: Session = Depends(get_db)):
    q = db.query(Employee).filter(Employee.employee_number == number.strip())
    if exclude_id is not None:
        q = q.filter(Employee.id != exclude_id)
    emp = q.first()
    return {"taken": emp is not None, "employee_name": emp.name if emp else None}


@router.get("/positions")
def list_positions(current_user: AppUser = Depends(_employee_list_access),
                   db: Session = Depends(get_db)):
    """Stillinger til medarbejder-modalens dropdown og registerets filter."""
    rows = db.query(MasterPosition).all()
    return [{"id": r.id, "name": r.name} for r in sorted(rows, key=lambda r: r.name.lower())]


@router.get("/anciennitet-alerts", response_model=list[AnciennitetsAlert])
def anciennitet_alerts(current_user: AppUser = Depends(require_permission("anciennitet_alert")),
                       db: Session = Depends(get_db)):
    """
    Medarbejdere der har opnået 9 måneders anciennitet, men hvor
    overenskomsttypen har en 9-mdr-variant, de endnu ikke er flyttet til.
    Springer medarbejdere over, hvor advarslen er afvist server-side.
    """
    alerts = []
    employees = db.query(Employee).filter(Employee.active == True).all()
    for emp in employees:
        months = _months_employed(emp.hire_date)
        if months < 9:
            continue
        if emp.anciennitet_dismissed_at is not None:
            continue
        variant = seniority_variant_exists_from_db(db, emp.agreement_type)
        if variant:
            alerts.append(AnciennitetsAlert(
                employee_id=emp.id,
                employee_name=emp.name,
                employee_number=emp.employee_number,
                hire_date=emp.hire_date,
                months_employed=months,
                suggested_agreement_type=variant,
            ))
    return alerts


@router.post("/{employee_id}/dismiss-anciennitet", status_code=204)
def dismiss_anciennitet(employee_id: int,
                        current_user: AppUser = Depends(require_permission("anciennitet_alert")),
                        db: Session = Depends(get_db)):
    """Marker anciennitetsadvarsel som afvist for denne medarbejder."""
    emp = db.query(Employee).filter(Employee.id == employee_id).first()
    if not emp:
        raise HTTPException(404, "Medarbejder ikke fundet")
    emp.anciennitet_dismissed_at = datetime.utcnow()
    db.commit()


@router.get("/paragraf56-alerts", response_model=Paragraf56AlertsResponse)
def paragraf56_alerts(current_user: AppUser = Depends(require_permission("paragraf_56_alert")),
                      db: Session = Depends(get_db)):
    """
    §56-advarsler for den aktuelle bruger: 'upcoming' (slutdato inden for 30 dage,
    §56 stadig aktiv) og 'expired' (§56 netop auto-deaktiveret pga. overskredet
    slutdato). Afvisning er pr. bruger (Paragraf56AlertDismissal), ikke global.
    """
    _sweep_expired_paragraf_56(db)
    today = date.today()
    window = today + timedelta(days=30)
    dismissed = {
        (d.employee_id, d.alert_type)
        for d in db.query(Paragraf56AlertDismissal).filter(
            Paragraf56AlertDismissal.user_id == current_user.id
        ).all()
    }
    upcoming = [
        _paragraf56_alert(e) for e in db.query(Employee).filter(
            Employee.paragraf_56 == True,
            Employee.paragraf_56_end_date.isnot(None),
            Employee.paragraf_56_end_date >= today,
            Employee.paragraf_56_end_date <= window,
        ).all()
        if (e.id, "upcoming") not in dismissed
    ]
    expired = [
        _paragraf56_alert(e) for e in db.query(Employee).filter(
            Employee.paragraf_56 == False,
            Employee.paragraf_56_end_date.isnot(None),
            Employee.paragraf_56_end_date < today,
        ).all()
        if (e.id, "expired") not in dismissed
    ]
    return Paragraf56AlertsResponse(upcoming=upcoming, expired=expired)


@router.post("/{employee_id}/dismiss-paragraf56-alert", status_code=204)
def dismiss_paragraf56_alert(employee_id: int, body: Paragraf56AlertDismiss,
                             current_user: AppUser = Depends(require_permission("paragraf_56_alert")),
                             db: Session = Depends(get_db)):
    """Marker en §56-advarsel som afvist for DEN AKTUELLE BRUGER (ikke globalt)."""
    if body.alert_type not in ("upcoming", "expired"):
        raise HTTPException(400, f"Ukendt alert_type: {body.alert_type}")
    existing = db.query(Paragraf56AlertDismissal).filter(
        Paragraf56AlertDismissal.employee_id == employee_id,
        Paragraf56AlertDismissal.user_id == current_user.id,
        Paragraf56AlertDismissal.alert_type == body.alert_type,
    ).first()
    if not existing:
        db.add(Paragraf56AlertDismissal(
            employee_id=employee_id, user_id=current_user.id, alert_type=body.alert_type
        ))
        db.commit()


_MILESTONE_PERMS = {"birthday": "birthday_alert", "jubilee": "jubilee_alert", "elev": "elev_alert"}


def _milestone_alerts_for(emp: Employee, today: date) -> list:
    """Alle aktuelle jubilæums-/elev-/fødselsdagsadvarsler for én medarbejder."""
    out = []
    if emp.cpr_number:
        try:
            hit = round_birthday_alert(cpr_birthdate(emp.cpr_number), today)
        except ValueError:
            hit = None
        if hit:
            age, when = hit
            out.append(("birthday", f"birthday_{age}", when, f"fylder {age} år"))
    hit = jubilee_alert(emp.seniority_date or emp.hire_date, today)
    if hit:
        years, when = hit
        out.append(("jubilee", f"jubilee_{years}", when, f"har {years} års jubilæum"))
    if (emp.agreement_kind != FUNKTIONAER and emp.elev and emp.elev_end_date
            and in_alert_window(emp.elev_end_date, today)):
        out.append(("elev", f"elev_{emp.elev_end_date.isoformat()}", emp.elev_end_date, "afslutter elevtiden"))
    return out


@router.get("/milestone-alerts", response_model=list[MilestoneAlert])
def milestone_alerts(current_user: AppUser = Depends(require_any_permission(*_MILESTONE_PERMS.values())),
                     db: Session = Depends(get_db), today: Optional[date] = None):
    """Jubilæum (25/40/50 år), elev slutter og rund fødselsdag – fra en måned før til og
    med dagen. Hver type kræver sin egen rettighed. Afvisning er pr. bruger."""
    today = today or date.today()
    allowed = {k for k, perm in _MILESTONE_PERMS.items() if user_has_permission(db, current_user, perm)}
    dismissed = {
        (d.employee_id, d.alert_type)
        for d in db.query(Paragraf56AlertDismissal).filter(
            Paragraf56AlertDismissal.user_id == current_user.id
        ).all()
    }
    alerts = []
    for emp in db.query(Employee).filter(Employee.active == True).all():
        for kind, key, when, label in _milestone_alerts_for(emp, today):
            if kind in allowed and (emp.id, key) not in dismissed:
                alerts.append(MilestoneAlert(
                    employee_id=emp.id, employee_name=emp.name, employee_number=emp.employee_number,
                    kind=kind, alert_key=key, event_date=when, label=label,
                ))
    return sorted(alerts, key=lambda a: (a.event_date, a.employee_name))


@router.post("/{employee_id}/dismiss-milestone-alert", status_code=204)
def dismiss_milestone_alert(employee_id: int, body: MilestoneAlertDismiss,
                            current_user: AppUser = Depends(require_any_permission(*_MILESTONE_PERMS.values())),
                            db: Session = Depends(get_db)):
    """'OK' i popup'en – advarslen kommer ikke igen for DENNE bruger og DENNE begivenhed."""
    kind = body.alert_key.split("_", 1)[0]
    if kind not in _MILESTONE_PERMS or "_" not in body.alert_key:
        raise HTTPException(400, f"Ukendt advarsel: {body.alert_key}")
    if not user_has_permission(db, current_user, _MILESTONE_PERMS[kind]):
        raise HTTPException(403, "Ingen adgang")
    exists = db.query(Paragraf56AlertDismissal).filter(
        Paragraf56AlertDismissal.employee_id == employee_id,
        Paragraf56AlertDismissal.user_id == current_user.id,
        Paragraf56AlertDismissal.alert_type == body.alert_key,
    ).first()
    if not exists:
        db.add(Paragraf56AlertDismissal(employee_id=employee_id, user_id=current_user.id,
                                        alert_type=body.alert_key))
        db.commit()


EXPORT_HEADERS = [
    "Lønnummer", "Navn", "Fuldlønnet", "Natarbejdetillæg", "Stilling", "Disponentgruppe",
    "Ansættelsesdato", "Telefon", "Mobil", "Email", "Elev", "Elev start", "Elev slut",
]


def _dk_date(d: Optional[date]) -> Optional[str]:
    return d.strftime("%d-%m-%Y") if d else None


@router.post("/export-xlsx")
def export_employees_xlsx(body: EmployeeExportRequest,
                          current_user: AppUser = Depends(require_permission("employee_export")),
                          db: Session = Depends(get_db)):
    """Medarbejderregisterets tabelvisning som Excel. Rækkefølgen er klientens
    (efter filtre/søgning/sortering). CPR kommer aldrig med."""
    by_id = {e.id: e for e in db.query(Employee).filter(Employee.id.in_(body.employee_ids)).all()}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medarbejderregister"
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="317423", end_color="317423", fill_type="solid")
    for emp_id in body.employee_ids:
        e = by_id.get(emp_id)
        if not e:
            continue
        is_elev = bool(e.elev) and e.agreement_kind != FUNKTIONAER
        ws.append([
            e.employee_number, e.name,
            "Ja" if e.fuldloennet else "Nej",
            "Ja" if e.natarbejde_tillaeg else "Nej",
            e.position.name if e.position else None,
            e.dispatcher_group.name if e.dispatcher_group else None,
            _dk_date(e.hire_date), e.phone, e.mobile, e.email,
            "Ja" if is_elev else "Nej",
            _dk_date(e.elev_start_date) if is_elev else None,
            _dk_date(e.elev_end_date) if is_elev else None,
        ])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, max(len(str(c.value or "")) for c in col) + 2)
    log_action(db, current_user, "employee_export", "employee", None,
               f"Eksporteret medarbejderregister ({len(by_id)} medarbejdere)")
    db.commit()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Medarbejderregister_{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{employee_id}", response_model=EmployeeResponse)
def get_employee(employee_id: int,
                 current_user: AppUser = Depends(_employee_list_access),
                 db: Session = Depends(get_db)):
    emp = db.query(Employee).filter(Employee.id == employee_id).first()
    if not emp:
        raise HTTPException(404, "Medarbejder ikke fundet")
    return _visible_response(emp, db, current_user)


_CLEARABLE_EMPLOYEE_FIELDS = (
    "tachograph_card_number", "initials", "address", "postal_code",
    "email", "phone", "mobile", "seniority_date", "elev_start_date", "elev_end_date",
    "position_id",
)


@router.patch("/{employee_id}", response_model=EmployeeResponse)
def update_employee(employee_id: int, body: EmployeeUpdate,
                    current_user: AppUser = Depends(require_permission("manage_employees")),
                    db: Session = Depends(get_db)):
    emp = db.query(Employee).filter(Employee.id == employee_id).first()
    if not emp:
        raise HTTPException(404, "Medarbejder ikke fundet")
    if body.employee_number and body.employee_number != emp.employee_number:
        if db.query(Employee).filter(Employee.employee_number == body.employee_number,
                                     Employee.id != emp.id).first():
            raise HTTPException(400, "Lønnummer eksisterer allerede")
    sent = set(body.model_fields_set)
    # CPR: maskeret værdi (fra en bruger uden 'Se CPR-nummer') betyder "uændret"
    cpr_sent = "cpr_number" in sent and not is_masked_cpr(body.cpr_number)
    new_cpr = _clean_cpr(body.cpr_number) if cpr_sent else emp.cpr_number
    sent.discard("cpr_number")
    effective = {
        c: (getattr(body, c) if c in sent else getattr(emp, c))
        for c in ("position_id", "email", "initials", "tachograph_card_number",
                  "elev", "elev_start_date", "elev_end_date")
    }
    _validate_employee_fields(db, body.agreement_kind or emp.agreement_kind, effective, sent)
    if body.agreement_kind and not db.query(MasterAgreementKind).filter(
        MasterAgreementKind.key == body.agreement_kind
    ).first():
        raise HTTPException(400, f"Ukendt aftaletype: {body.agreement_kind}")
    effective_kind = body.agreement_kind or emp.agreement_kind
    effective_agreement_type = (
        body.agreement_type if body.agreement_type is not None else emp.agreement_type
    )
    if _agreement_type_required(db, effective_kind):
        if not effective_agreement_type or effective_agreement_type not in load_agreement_types_from_db(db):
            raise HTTPException(400, f"Ukendt overenskomsttype: {effective_agreement_type}")
    elif body.agreement_type is None and body.agreement_kind and body.agreement_kind != emp.agreement_kind:
        # Skiftes til en type der ikke kræver Overenskomsttype, uden at et nyt
        # felt er angivet samtidig – nulstil det gemte felt til "ikke relevant".
        body.agreement_type = ""
    old_agreement_type = emp.agreement_type
    _paragraf56_excludes = {"dispatcher_group_id", "fast_bil_vehicle_id", "absence_vehicle_id",
                            "paragraf_56", "paragraf_56_start_date", "paragraf_56_end_date",
                            "cpr_number"}
    for field_name, value in body.model_dump(exclude_none=True, exclude=_paragraf56_excludes).items():
        if field_name == "work_schedule":
            value = body.work_schedule.model_dump()
        setattr(emp, field_name, value)
    # Valgfri tekstfelter der sendes som tomme (null) skal tømmes – ikke bevare den
    # gamle værdi, som exclude_none ovenfor ellers ville gøre.
    for field_name in _CLEARABLE_EMPLOYEE_FIELDS:
        if field_name in body.model_fields_set and getattr(body, field_name) is None:
            setattr(emp, field_name, None)
    if cpr_sent:
        emp.cpr_number = new_cpr
    if "dispatcher_group_id" in body.model_fields_set:
        emp.dispatcher_group = _resolve_dispatcher_group(db, body.dispatcher_group_id)
    if "fast_bil_vehicle_id" in body.model_fields_set:
        emp.fast_bil_vehicle_id = _resolve_vehicle_id(db, body.fast_bil_vehicle_id)
    if "absence_vehicle_id" in body.model_fields_set:
        emp.absence_vehicle_id = _resolve_vehicle_id(db, body.absence_vehicle_id)
    if "paragraf_56" in body.model_fields_set:
        start, end = _validate_paragraf_56(
            bool(body.paragraf_56), body.paragraf_56_start_date, body.paragraf_56_end_date
        )
        if end != emp.paragraf_56_end_date:
            db.query(Paragraf56AlertDismissal).filter(
                Paragraf56AlertDismissal.employee_id == emp.id,
                Paragraf56AlertDismissal.alert_type.in_(("upcoming", "expired")),
            ).delete(synchronize_session=False)
        emp.paragraf_56 = bool(body.paragraf_56)
        emp.paragraf_56_start_date = start
        emp.paragraf_56_end_date = end
    # Nulstil afvist anciennitetsadvarsel hvis overenskomsttype er ændret
    if body.agreement_type and body.agreement_type != old_agreement_type:
        emp.anciennitet_dismissed_at = None
    db.commit()
    db.refresh(emp)
    return _apply_cpr_mask(_to_response(emp, db), db, current_user)
